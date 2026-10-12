#!/usr/bin/env python3
"""Make a Word (.docx) copy of a CV from its .tex source.

Some job sites (older Taleo and iCIMS forms, some local job boards) read Word
files more reliably than PDF. The .docx has the same text and order as the PDF:
one column, real Word headings, bullet lists, no tables, text boxes or images.

Usage: python3 tex2docx.py cv-systems-rust.tex [more.tex ...]
Needs: pip install python-docx
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

NAVY = RGBColor(0x1B, 0x2A, 0x47)
GRAY = RGBColor(0x4A, 0x55, 0x68)
LINK = RGBColor(0x1F, 0x5F, 0xA8)
FONT = "Calibri"
# Wide en-space gaps around separators, like \\enspace in LaTeX, so parsers see
# the job title, company and dates as separate pieces of text.
SEP = "\u2002|\u2002"

# Contact details, the same as \cvheader in cv-style.sty.
EMAIL = "abdelrahman.khalifa.gad@gmail.com"
PHONE = "+20 114 250 5042"
CITY = "Cairo, Egypt"
LINKS = [
    ("https://www.linkedin.com/in/abd002", "linkedin.com/in/abd002"),
    ("https://github.com/Abd002", "github.com/Abd002"),
    ("https://codeforces.com/profile/_Khalifa", "codeforces.com/profile/_Khalifa"),
]


def plain(s):
    """LaTeX escapes and spacing commands to plain Unicode text."""
    for a, b in [
        (r"\enspace|\enspace", SEP), (r"\enspace", "\u2002"), ("$\\cdot$", "·"),
        (r"\textasciitilde", "~"), (r"\%", "%"), (r"\#", "#"), (r"\_", "_"),
        (r"\&", "&"), (r"\$", "$"), ("---", "—"), ("--", "–"), ("~", " "),
        (r"\\", " "), ("\\,", " "),
    ]:
        s = s.replace(a, b)
    return s


def args(s, i, n):
    """Read n brace groups starting at s[i]; return (list, end index)."""
    out = []
    for _ in range(n):
        while s[i] in " \n":
            i += 1
        assert s[i] == "{", s[i:i + 40]
        depth, j = 0, i
        while True:
            if s[j] == "\\":
                j += 2
                continue
            if s[j] == "{":
                depth += 1
            elif s[j] == "}":
                depth -= 1
                if depth == 0:
                    break
            j += 1
        out.append(s[i + 1:j])
        i = j + 1
    return out, i


def add_link(par, url, text, size):
    r_id = par.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FA8")
    rpr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.append(sz)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    par._p.append(link)


def rich(par, s, size=10, bold=False, italic=False, color=None):
    """Write LaTeX inline text (with \\textbf, \\texttt, \\cvlink) into a paragraph."""
    pos = 0
    pattern = re.compile(r"\\(textbf|texttt|textit|cvlink)\{")
    while True:
        m = pattern.search(s, pos)
        if not m:
            break
        if m.start() > pos:
            run(par, plain(s[pos:m.start()]), size, bold, italic, color)
        cmd = m.group(1)
        n = 2 if cmd == "cvlink" else 1
        a, pos = args(s, m.end() - 1, n)
        if cmd == "cvlink":
            add_link(par, a[0], plain(re.sub(r"\\\w+\{|\}", "", a[1])), size)
        else:
            rich(par, a[0], size, bold or cmd == "textbf", italic or cmd == "textit", color)
    if pos < len(s):
        run(par, plain(s[pos:]), size, bold, italic, color)


def run(par, text, size, bold, italic, color):
    text = re.sub(r"[ \n]+", " ", text)
    if not text:
        return
    r = par.add_run(text)
    r.font.size = Pt(size)
    r.font.name = FONT
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color


def para(doc, style=None, space_before=0, align=None):
    p = doc.add_paragraph(style=style)
    f = p.paragraph_format
    f.space_before = Pt(space_before)
    f.space_after = Pt(0)
    f.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    return p


def convert(tex_path):
    s = open(tex_path, encoding="utf-8").read()
    s = re.sub(r"(?m)(?<!\\)%.*$", "", s)
    s = s[s.index(r"\begin{document}") + len(r"\begin{document}"):s.index(r"\end{document}")]

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)  # A4
    sec.left_margin = sec.right_margin = Inches(0.5)
    sec.top_margin = sec.bottom_margin = Inches(0.4)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT
    h1.font.size = Pt(12)
    h1.font.bold = True
    h1.font.color.rgb = NAVY
    h1.paragraph_format.space_before = Pt(4)
    h1.paragraph_format.space_after = Pt(2)
    h1.paragraph_format.keep_with_next = True

    i = 0
    in_list = False
    item = None
    token = re.compile(r"\\(cvheader|section\*|cvjob|cventry|cvsub|cvskill|begin\{itemize\}|end\{itemize\}|item)(?![A-Za-z])")
    while i < len(s):
        m = token.search(s, i)
        text_before = s[i:m.start()] if m else s[i:]
        if text_before.strip():
            if in_list and item is not None:
                rich(item, " " + text_before.strip())
            else:
                rich(para(doc), text_before.strip())
        if not m:
            break
        cmd = m.group(1)
        i = m.end()
        if cmd == "cvheader":
            (name, title), i = args(s, i, 2)
            rich(para(doc, align=WD_ALIGN_PARAGRAPH.CENTER), name, size=20, bold=True, color=NAVY)
            rich(para(doc, align=WD_ALIGN_PARAGRAPH.CENTER), title, size=12, color=NAVY)
            p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
            run(p, f"{EMAIL}{SEP}{PHONE}{SEP}{CITY}", 9.5, False, False, GRAY)
            p = para(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
            for k, (url, text) in enumerate(LINKS):
                if k:
                    run(p, SEP, 9.5, False, False, GRAY)
                add_link(p, url, text, 9.5)
        elif cmd == "section*":
            (title,), i = args(s, i, 1)
            doc.add_heading(plain(title).upper(), level=1)
        elif cmd == "cvjob":
            (title, company, dates), i = args(s, i, 3)
            p = para(doc, space_before=3)
            p.paragraph_format.keep_with_next = True
            rich(p, title, bold=True, color=NAVY)
            run(p, SEP, 10, False, False, GRAY)
            rich(p, company, bold=True, color=NAVY)
            run(p, SEP, 10, False, False, GRAY)
            rich(p, dates, size=9.5, bold=True, color=GRAY)
        elif cmd == "cventry":
            (title, dates), i = args(s, i, 2)
            p = para(doc, space_before=3)
            p.paragraph_format.keep_with_next = True
            rich(p, title, bold=True, color=NAVY)
            run(p, SEP, 10, False, False, GRAY)
            rich(p, dates, size=9.5, bold=True, color=GRAY)
        elif cmd == "cvsub":
            (text,), i = args(s, i, 1)
            p = para(doc)
            p.paragraph_format.keep_with_next = True
            rich(p, text, size=9.5, italic=True, color=GRAY)
        elif cmd == "cvskill":
            (label, items), i = args(s, i, 2)
            p = para(doc)
            rich(p, label + ": ", bold=True)
            rich(p, items)
        elif cmd == "begin{itemize}":
            in_list, item = True, None
        elif cmd == "end{itemize}":
            in_list, item = False, None
        elif cmd == "item":
            # A plain "•" and a tab, not Word list numbering: every parser
            # knows this bullet character (Word's Symbol-font bullet is not).
            item = para(doc)
            item.paragraph_format.left_indent = Inches(0.18)
            item.paragraph_format.first_line_indent = Inches(-0.13)
            item.paragraph_format.tab_stops.add_tab_stop(Inches(0.18))
            run(item, "•\t", 10, False, False, None)
            j = s.find("\\item", i)
            k = s.find("\\end{itemize}", i)
            end = min(x for x in (j, k) if x != -1)
            rich(item, s[i:end].strip())
            i = end
    out = tex_path[:-4] + ".docx"
    doc.core_properties.author = "Abdelrahman Khalifa"
    doc.core_properties.title = "Abdelrahman Khalifa - CV"
    doc.save(out)
    return out


if __name__ == "__main__":
    for path in sys.argv[1:]:
        print(convert(path))
